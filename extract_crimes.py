"""
CrimeTime data extraction tool.

Connects to the restored CrimeTime SQL Server database, pulls every active
crime with its metadata and sentencing-coverage status, and outputs an Excel
workbook for review.

Requirements:
    pip install pymssql pandas openpyxl

Usage:
    python extract_crimes.py

Adjust the connection string below to match your local SQL Server setup.
"""

import pymssql
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ---- Connection config -----------------------------------------------
DB_HOST   = "localhost"
DB_PORT   = 1433
DB_USER   = "sa"
DB_PASS   = "YourStrong!Passw0rd"
DB_NAME   = "CrimeTime"

OUTPUT_PATH = "crimetime_extraction.xlsx"


def extract():
    query = """
    SELECT
        c.CrimesID,
        c.section,
        c.subsection,
        c.full_law_description,
        c.class,
        lc.Literal AS ClassLiteral,
        c.deprecated,
        c.effective_date,
        c.repeal_date,
        (SELECT COUNT(*) FROM CrimeAnswer ca WHERE ca.CrimesID = c.CrimesID) AS ScenarioCount,
        (SELECT COUNT(DISTINCT al.AnswerID)
         FROM CrimeAnswer ca
         JOIN AnswerLink al ON al.CrimeAnswerID = ca.CrimeAnswerID
         WHERE ca.CrimesID = c.CrimesID) AS DistinctAnswerCount
    FROM Crimes c
    LEFT JOIN LawClass lc ON c.class = lc.LawClassID
    WHERE c.deprecated = 0
    ORDER BY c.section, c.subsection;
    """
    with pymssql.connect(
        server=DB_HOST, port=DB_PORT,
        user=DB_USER, password=DB_PASS,
        database=DB_NAME, as_dict=False,
    ) as conn:
        df = pd.read_sql(query, conn)
    return df


def write_excel(df: pd.DataFrame):
    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Crimes", index=False)

    style_workbook()


def style_workbook():
    from openpyxl import load_workbook

    wb = load_workbook(OUTPUT_PATH)
    header_fill = PatternFill("solid", start_color="2E4A7D", end_color="2E4A7D")
    header_font = Font(bold=True, color="FFFFFF")

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for col in sheet.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 60)
        sheet.freeze_panes = "A2"

    wb.save(OUTPUT_PATH)


def main():
    print("Extracting crimes from CrimeTime database...")
    df = extract()
    print(f"  Pulled {len(df)} active crimes.")

    print(f"Writing {OUTPUT_PATH}...")
    write_excel(df)
    print("Done.")


if __name__ == "__main__":
    main()